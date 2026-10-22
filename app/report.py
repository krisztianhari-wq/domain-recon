"""Exportformátumok: olvasható Markdown-jelentés és tiszta, hosztonként csoportosított JSON.

Cél: az egy hoszthoz tartozó adatok (DNS, hálózat/ASN, portok, technológia, kitettség, sebezhetőség,
takeover) jól elkülönítve, átlátható szekciókban jelenjenek meg — ne egy lapos, ismétlődő táblában.
"""
from __future__ import annotations

import html
import io
import time

E = html.escape


def _status(h: dict) -> str:
    if h.get("takeover") and h["takeover"].get("service"):
        return "takeover"
    if (h.get("exposed_flags") or h.get("vuln") or h.get("cve")):
        return "exposed"
    return "live" if h.get("alive") else "dead"


def _fmt_dt(epoch) -> str:
    try:
        return time.strftime("%Y-%m-%d %H:%M", time.localtime(epoch))
    except Exception:  # noqa: BLE001
        return "-"


def _host_clean(h: dict) -> dict:
    """Egy host adatai tiszta, csoportosított szerkezetben (JSON-exporthoz)."""
    net = []
    for i in h.get("ips_info", []) or [{"ip": ip} for ip in h.get("ips", [])]:
        net.append({k: i.get(k) for k in ("ip", "ptr", "asn", "org", "cc", "prefix") if i.get(k) is not None})
    ports = []
    for p in h.get("ports", []):
        rec = {"port": p["port"], "service": p.get("name"), "expected": p.get("expected", False)}
        if p.get("rtt_ms") is not None:
            rec["rtt_ms"] = p["rtt_ms"]
        if p.get("http"):
            http = p["http"]
            rec["http"] = {k: http[k] for k in ("status", "server", "powered_by", "title") if http.get(k)}
        if p.get("tls"):
            tls = p["tls"]
            rec["tls"] = {k: tls[k] for k in ("cn", "sans", "expires") if tls.get(k)}
        if p.get("banner"):
            rec["banner"] = p["banner"]
        if p.get("tech"):
            rec["tech"] = p["tech"]
        ports.append(rec)
    out = {
        "host": h["host"],
        "status": _status(h),
        "dns": {
            "a": h.get("v4", []), "aaaa": h.get("v6", []),
            "cname": h.get("cname"), "wildcard": bool(h.get("wildcard")),
            "source": "brute" if h.get("source_brute") else "passive",
        },
        "network": net,
        "ports": ports,
    }
    if h.get("tech"):
        out["technologies"] = h["tech"]
    if h.get("exposed_flags"):
        out["exposure"] = h["exposed_flags"]
    if h.get("vuln"):
        out["vulnerabilities"] = h["vuln"]
    if h.get("cve"):
        out["known_cves"] = h["cve"]
    if h.get("takeover") and h["takeover"].get("service"):
        out["takeover"] = h["takeover"]
    if h.get("pq") and h["pq"].get("kex"):
        pq = h["pq"]
        out["quantum"] = {"verdict": pq.get("verdict"), "key_exchange": pq.get("kex"), "certificate": pq.get("cert")}
    if h.get("tls_sans"):
        out["tls_sans"] = h["tls_sans"]
    return out


def build_json(scan: dict, hosts: list[dict]) -> dict:
    """Tiszta, csoportosított JSON: scan-meta + posture + hostonkénti blokkok."""
    return {
        "scan": {
            "domain": scan.get("domain"),
            "id": scan.get("id"),
            "created": _fmt_dt(scan.get("created")),
            "finished": _fmt_dt(scan.get("finished")) if scan.get("finished") else None,
            "version": scan.get("version"),
            "build": scan.get("build"),
            "passive_status": scan.get("status"),
            "active_status": scan.get("active_status"),
            "counts": {"hosts": scan.get("n_hosts", 0), "live": scan.get("n_live", 0),
                       "open_ports": scan.get("n_ports", 0)},
        },
        "dns_posture": scan.get("posture"),
        "wildcard_ips": scan.get("wildcard") or [],
        "hosts": [_host_clean(h) for h in hosts],
    }


# ---------------- Markdown-jelentés ----------------
def _net_line(i: dict) -> str:
    parts = [i.get("ip", "")]
    if i.get("ptr"):
        parts.append(f"PTR {i['ptr']}")
    if i.get("asn"):
        parts.append(f"AS{i['asn']}" + (f" {i['org']}" if i.get("org") else "") + (f" ({i['cc']})" if i.get("cc") else ""))
    return " · ".join(p for p in parts if p)


def _host_md(h: dict) -> str:
    st = _status(h)
    badge = {"takeover": "🔴 TAKEOVER", "exposed": "🟠 KITETT", "live": "🟢 élő", "dead": "⚪ nem él"}[st]
    lines = [f"### {h['host']}  —  {badge}", ""]
    # DNS
    dns = []
    if h.get("ips"):
        dns.append(f"**IP:** {', '.join(h['ips'])}")
    if h.get("cname"):
        dns.append(f"**CNAME:** {h['cname']}")
    dns.append("**Forrás:** " + ("brute" if h.get("source_brute") else "passzív") + (" · wildcard" if h.get("wildcard") else ""))
    lines.append(" · ".join(dns))
    # hálózat/ASN
    info = h.get("ips_info") or []
    if any(i.get("asn") or i.get("ptr") for i in info):
        lines.append("")
        lines.append("**Hálózat:**")
        for i in info:
            line = _net_line(i)
            if line:
                lines.append(f"- {line}")
    # technológia
    if h.get("tech"):
        lines.append("")
        lines.append("**Technológia:** " + ", ".join(h["tech"]))
    # kvantum-reziliencia
    if h.get("pq") and h["pq"].get("kex"):
        pq = h["pq"]
        k, c = pq.get("kex", {}), pq.get("cert", {})
        vmap = {"hybrid": "kvantum-biztos (hibrid ML-KEM)", "classical": "klasszikus — HNDL-kockázat", "unknown": "ismeretlen"}
        det = []
        if k.get("group"):
            det.append(f"kulcscsere {k['group']}")
        if k.get("tls"):
            det.append(f"TLS {k['tls']}")
        if c.get("key_alg"):
            det.append(f"cert {c['key_alg']}" + (f" {c['key_bits']}" if c.get("key_bits") else ""))
        if c.get("sig_alg"):
            det.append(f"aláírás {c['sig_alg']}")
        lines.append("")
        lines.append(f"**Kvantum-reziliencia (TLS):** {vmap.get(pq.get('verdict'), '?')}"
                     + (" · " + " · ".join(det) if det else ""))
    # takeover
    if h.get("takeover") and h["takeover"].get("service"):
        t = h["takeover"]
        lines.append("")
        lines.append(f"**Lehetséges takeover:** {t['service']} → {t.get('cname','')} "
                     f"({'megerősítve' if t.get('confirmed') else 'gyanús'})")
    # kitettség
    if h.get("exposed_flags"):
        lines.append("")
        lines.append("**Kitett szolgáltatás:** " + ", ".join(h["exposed_flags"]))
    # sebezhetőség
    if h.get("vuln"):
        lines.append("")
        lines.append("**Sebezhetőség / kitettség:**")
        for v in h["vuln"]:
            ev = f" — {v['evidence']}" if v.get("evidence") else ""
            lines.append(f"- [{v.get('severity','?')}] {v.get('title','')}{ev}")
    # ismert CVE (verzió alapján)
    if h.get("cve"):
        lines.append("")
        lines.append("**Ismert sebezhetőség (CVE, verzió alapján):**")
        for v in h["cve"]:
            ev = f" — {v['evidence']}" if v.get("evidence") else ""
            lines.append(f"- [{v.get('severity','?')}] {v.get('title','')}{ev}")
    # portok
    if h.get("ports"):
        lines.append("")
        lines.append("**Nyitott portok:**")
        lines.append("")
        lines.append("| Port | Szolgáltatás | Részlet | RTT |")
        lines.append("|---|---|---|---|")
        for p in h["ports"]:
            http = p.get("http") or {}
            tls = p.get("tls") or {}
            det = []
            if http.get("status"):
                det.append(f"HTTP {http['status']}")
            if http.get("server"):
                det.append(http["server"])
            if http.get("powered_by"):
                det.append(http["powered_by"])
            if http.get("title"):
                det.append(f"„{http['title']}”")
            if p.get("tech"):
                det.append(", ".join(p["tech"]))
            if tls.get("cn"):
                det.append(f"CN={tls['cn']}")
            if p.get("banner"):
                det.append(p["banner"])
            seen, ded = set(), []
            for d in det:
                d = str(d)
                if d.lower() not in seen:
                    seen.add(d.lower())
                    ded.append(d)
            detail = "; ".join(ded).replace("|", "\\|") or "—"
            rtt = f"{p['rtt_ms']} ms" if p.get("rtt_ms") is not None else "—"
            lines.append(f"| {p['port']} | {p.get('name','?')} | {detail} | {rtt} |")
    if h.get("tls_sans"):
        lines.append("")
        lines.append("**TLS SAN (új nevek):** " + ", ".join(h["tls_sans"]))
    lines.append("")
    return "\n".join(lines)


def build_markdown(scan: dict, hosts: list[dict]) -> str:
    domain = scan.get("domain", "")
    p = scan.get("posture") or {}
    wc = scan.get("wildcard") or []
    n_live = sum(1 for h in hosts if h.get("alive"))
    n_exp = sum(1 for h in hosts if _status(h) == "exposed")
    n_tk = sum(1 for h in hosts if _status(h) == "takeover")
    n_ports = sum(len(h.get("ports", [])) for h in hosts)

    L = [f"# domain-recon — {domain}", ""]
    ver = f" · v{scan['version']}" + (f" ({scan['build']})" if scan.get("build") else "") if scan.get("version") else ""
    L.append(f"*Vizsgálat: {_fmt_dt(scan.get('created'))}{ver} · "
             f"passzív: {scan.get('status','?')} · aktív: {scan.get('active_status','none')}*")
    L.append("")
    L.append(f"**Összegzés:** {len(hosts)} host · {n_live} élő · {n_tk} takeover · {n_exp} kitett · {n_ports} nyitott port")
    L.append("")
    # posture
    L.append("## DNS / e-mail posture (apex)")
    def ok(v): return "✓ van" if v else "✗ hiányzik"
    L.append("")
    L.append(f"- **SPF:** {ok(p.get('spf_ok'))}" + (f" — `{p['spf']}`" if p.get("spf") else ""))
    L.append(f"- **DMARC:** {ok(p.get('dmarc_ok'))}" + (f" — `{p['dmarc']}`" if p.get("dmarc") else ""))
    L.append(f"- **DKIM:** " + (", ".join(p["dkim"]) if p.get("dkim") else "✗ nincs ismert selector"))
    L.append(f"- **CAA:** " + (", ".join(p["caa"]) if p.get("caa") else "✗ nincs"))
    L.append(f"- **MX:** " + (", ".join(p["mx"]) if p.get("mx") else "✗ nincs"))
    L.append(f"- **NS:** " + (", ".join(p["ns"]) if p.get("ns") else "—"))
    if wc:
        L.append(f"- **Wildcard DNS:** ⚠ aktív — {', '.join(wc)}")
    L.append("")

    # hostok három csoportban: figyelmet igénylő / élő / nem élő
    order = {"takeover": 0, "exposed": 1, "live": 2, "dead": 3}
    attn = sorted([h for h in hosts if _status(h) in ("takeover", "exposed")],
                  key=lambda h: (order[_status(h)], h["host"]))
    live = sorted([h for h in hosts if _status(h) == "live"], key=lambda h: h["host"])
    dead = sorted([h for h in hosts if _status(h) == "dead"], key=lambda h: h["host"])

    if attn:
        L.append("## ⚠ Figyelmet igénylő hostok")
        L.append("")
        for h in attn:
            L.append(_host_md(h))
            L.append("---")
            L.append("")
    if live:
        L.append("## Élő hostok")
        L.append("")
        for h in live:
            L.append(_host_md(h))
            L.append("---")
            L.append("")
    if dead:
        L.append("## Nem élő / fel nem oldott hostok")
        L.append("")
        for h in dead:
            cn = f" (CNAME {h['cname']})" if h.get("cname") else ""
            L.append(f"- {h['host']}{cn}")
        L.append("")
    L.append("---")
    L.append("*sadrobot · domain-recon · csak engedélyezett célpontra*")
    return "\n".join(L)


# ---------------- XLSX (formázott, hosztonként csoportosítva; sadrobot paletta, nem Yettel) ----------------
# Semleges paletta: sötét kékeszöld fejléc, cián akcentcsík, halvány host-sávozás.
_DARK, _ACCENT, _PALE = "0E2733", "0B8FA0", "E7F2F6"
_X_HEADERS = ["Host", "Alive", "IP Address(es)", "CNAME", "Port", "Service",
              "HTTP Status", "Server", "Page Title", "TLS CN", "Banner"]
_X_WIDTHS = [39, 8, 49, 39, 8, 14, 13, 34, 34, 32, 37]
_X_CENTER = {2, 5, 7}
_X_NUM = {5, 7}
_X_STATUS = [(200, 299, "D6EEDD", "1F8A4C"), (300, 399, "DCEBF3", "1F5E80"),
             (400, 499, "FBE6CF", "9A6A12"), (500, 599, "F6D9D5", "B02A1E")]


def _x_rows(hosts: list[dict]) -> list[list]:
    """Egy sor portonként; a host-mezők MINDEN során kitöltve (a hostonkénti sávozás így működik)."""
    order = {"takeover": 0, "exposed": 1, "live": 2, "dead": 3}
    out = []
    for h in sorted(hosts, key=lambda x: (order[_status(x)], x["host"])):
        alive = "yes" if h.get("alive") else "no"
        ips = " ".join(h.get("ips", []))
        cname = h.get("cname") or ""
        ports = h.get("ports") or []
        if not ports:
            out.append([h["host"], alive, ips, cname, "", "", "", "", "", "", ""])
            continue
        for p in ports:
            http = p.get("http") or {}
            tls = p.get("tls") or {}
            out.append([h["host"], alive, ips, cname, p["port"], p.get("name", ""),
                        http.get("status", ""), http.get("server", ""), http.get("title", ""),
                        tls.get("cn", ""), (p.get("banner") or "")[:200]])
    return out


def build_xlsx(scan: dict, hosts: list[dict], classification: str = "Confidential") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.styles.differential import DifferentialStyle
    from openpyxl.formatting.rule import Rule, FormulaRule, CellIsRule
    from openpyxl.comments import Comment

    def fill(c):
        return PatternFill(start_color=c, end_color=c, fill_type="solid")

    domain = scan.get("domain", "")
    scan_id = scan.get("id", "")
    ts = _fmt_dt(scan.get("created"))
    data = _x_rows(hosts)
    first, last = 5, max(5, 4 + len(data))

    wb = Workbook()
    ws = wb.active
    ws.title = f"recon_{domain}_{scan_id}"[:31]
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = _DARK

    ws["A1"] = f"Domain Recon — {domain}"
    ws["A1"].font = Font(name="Calibri", size=18, bold=True, color=_DARK)
    ws["K1"] = classification
    ws["K1"].font = Font(name="Calibri", size=10, color="000000")
    ws["K1"].alignment = Alignment(horizontal="right")
    ws["A2"] = ('="External attack-surface scan   ·   "&COUNTA($A$5:$A$%d)&'
                '"   records   ·   Scan %s   ·   ID %s"') % (last, ts, scan_id)
    ws["A2"].font = Font(name="Calibri", size=10, color=_DARK)
    ws["A2"].comment = Comment(f"domain-recon scan · domain {domain} · id {scan_id} · {ts} · "
                               f"v{scan.get('version','')} {scan.get('build','')}".strip(), "recon")
    for c in range(1, 12):
        ws.cell(3, c).fill = fill(_ACCENT)
    ws.row_dimensions[3].height = 6

    for c, h in enumerate(_X_HEADERS, 1):
        cell = ws.cell(4, c, h)
        cell.fill = fill(_DARK)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[4].height = 22

    for r, row in enumerate(data, first):
        for c in range(1, 12):
            v = row[c - 1] if c - 1 < len(row) else ""
            v = None if v == "" else v
            if v is not None and c in _X_NUM:
                try:
                    v = int(v)
                except (ValueError, TypeError):
                    pass
            cell = ws.cell(r, c, v)
            cell.font = Font(name="Calibri", size=11)
            cell.alignment = Alignment(vertical="center",
                                       horizontal="center" if c in _X_CENTER else None)

    for i, w in enumerate(_X_WIDTHS):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:K{last}"

    body, gcol = f"A{first}:K{last}", f"G{first}:G{last}"
    prio = 1
    for lo, hi, bg, fc in _X_STATUS:
        rule = CellIsRule(operator="between", formula=[str(lo), str(hi)], fill=fill(bg), font=Font(color=fc))
        rule.priority = prio
        prio += 1
        ws.conditional_formatting.add(gcol, rule)
    divider = Rule(type="expression", formula=["$A5<>$A4"],
                   dxf=DifferentialStyle(border=Border(top=Side(style="thin", color=_DARK))))
    divider.priority = prio
    prio += 1
    ws.conditional_formatting.add(body, divider)
    band = FormulaRule(formula=["ISODD(SUMPRODUCT(--($A$5:$A5<>$A$4:$A4)))"], fill=fill(_PALE))
    band.priority = prio
    ws.conditional_formatting.add(body, band)

    ws.oddHeader.right.text = classification
    ws.oddFooter.right.text = "&P"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------------- Nyomtatható HTML-jelentés (böngészőből PDF) ----------------
_STATUS_BADGE = {"takeover": ("TAKEOVER", "bad"), "exposed": ("KITETT", "warn"),
                 "live": ("élő", "good"), "dead": ("nem él", "muted")}


def _host_html(h: dict) -> str:
    st = _status(h)
    label, cls = _STATUS_BADGE[st]
    out = [f'<section class="host {cls}"><h3>{E(h["host"])} <span class="badge {cls}">{label}</span></h3>']
    meta = []
    if h.get("ips"):
        meta.append("IP: " + E(", ".join(h["ips"])))
    if h.get("cname"):
        meta.append("CNAME: " + E(h["cname"]))
    meta.append("Forrás: " + ("brute" if h.get("source_brute") else "passzív"))
    out.append('<p class="meta">' + " · ".join(meta) + "</p>")
    info = h.get("ips_info") or []
    if any(i.get("asn") or i.get("ptr") for i in info):
        rows = "".join(f"<tr><td>{E(i.get('ip',''))}</td><td>{E(i.get('ptr') or '—')}</td>"
                       f"<td>{('AS'+E(str(i['asn']))) if i.get('asn') else '—'}</td>"
                       f"<td>{E(i.get('org') or '—')}</td><td>{E(i.get('cc') or '—')}</td></tr>" for i in info)
        out.append('<table class="t"><tr><th>IP</th><th>PTR</th><th>ASN</th><th>Org</th><th>CC</th></tr>' + rows + "</table>")
    if h.get("tech"):
        out.append('<p><b>Technológia:</b> ' + E(", ".join(h["tech"])) + "</p>")
    if h.get("takeover") and h["takeover"].get("service"):
        t = h["takeover"]
        out.append(f'<p class="warn"><b>Lehetséges takeover:</b> {E(t["service"])} → {E(t.get("cname",""))} '
                   f'({"megerősítve" if t.get("confirmed") else "gyanús"})</p>')
    for key, title in (("cve", "Ismert sebezhetőség (CVE)"), ("vuln", "Kitettség / sebezhetőség")):
        if h.get(key):
            items = "".join(f'<li class="sev-{E(str(v.get("severity","")))}"><b>{E(str(v.get("severity","")))}:</b> '
                            f'{E(v.get("title",""))}{(" — " + E(v["evidence"])) if v.get("evidence") else ""}</li>'
                            for v in h[key])
            out.append(f"<p><b>{title}:</b></p><ul class=\"find\">{items}</ul>")
    if h.get("exposed_flags"):
        out.append('<p><b>Kitett szolgáltatás:</b> ' + E(", ".join(h["exposed_flags"])) + "</p>")
    if h.get("pq") and h["pq"].get("kex"):
        pq = h["pq"]; k = pq.get("kex", {}); c = pq.get("cert", {})
        vmap = {"hybrid": "kvantum-biztos (hibrid ML-KEM)", "classical": "klasszikus — HNDL-kockázat", "unknown": "ismeretlen"}
        det = []
        if k.get("group"):
            det.append("kulcscsere " + E(k["group"]))
        if k.get("tls"):
            det.append("TLS " + E(str(k["tls"])))
        if c.get("key_alg"):
            det.append("cert " + E(c["key_alg"]) + (f" {c['key_bits']}" if c.get("key_bits") else ""))
        if c.get("sig_alg"):
            det.append("aláírás " + E(c["sig_alg"]))
        out.append(f'<p><b>Kvantum-reziliencia (TLS):</b> {vmap.get(pq.get("verdict"),"?")}'
                   + (" · " + " · ".join(det) if det else "") + "</p>")
    if h.get("ports"):
        rows = ""
        for p in h["ports"]:
            http = p.get("http") or {}; tls = p.get("tls") or {}
            det = []
            if http.get("status"):
                det.append("HTTP " + str(http["status"]))
            for kk in ("server", "powered_by", "title"):
                if http.get(kk):
                    det.append(E(str(http[kk])))
            if p.get("tech"):
                det.append(E(", ".join(p["tech"])))
            if tls.get("cn"):
                det.append("CN=" + E(tls["cn"]))
            if p.get("banner"):
                det.append(E(p["banner"]))
            rows += (f"<tr><td>{p['port']}</td><td>{E(p.get('name','?'))}</td>"
                     f"<td>{' · '.join(str(d) for d in det) or '—'}</td>"
                     f"<td>{(str(p['rtt_ms'])+' ms') if p.get('rtt_ms') is not None else '—'}</td></tr>")
        out.append('<table class="t"><tr><th>Port</th><th>Szolgáltatás</th><th>Részlet</th><th>RTT</th></tr>' + rows + "</table>")
    if h.get("tls_sans"):
        out.append('<p><b>TLS SAN:</b> ' + E(", ".join(h["tls_sans"])) + "</p>")
    out.append("</section>")
    return "".join(out)


def build_html(scan: dict, hosts: list[dict]) -> str:
    domain = scan.get("domain", "")
    p = scan.get("posture") or {}
    wc = scan.get("wildcard") or []
    n_live = sum(1 for h in hosts if h.get("alive"))
    n_exp = sum(1 for h in hosts if _status(h) == "exposed")
    n_tk = sum(1 for h in hosts if _status(h) == "takeover")
    n_cve = sum(1 for h in hosts if h.get("cve"))
    n_ports = sum(len(h.get("ports", [])) for h in hosts)
    ver = (f"v{scan['version']}" + (f" ({scan['build']})" if scan.get("build") else "")) if scan.get("version") else ""

    def stat(lbl, val):
        return f'<div class="s"><b>{lbl}</b><span>{val}</span></div>'
    stats = stat("Host", len(hosts)) + stat("Élő", n_live) + stat("Takeover", n_tk) + stat("Kitett", n_exp) \
        + stat("Nyitott port", n_ports) + (stat("CVE-találat", n_cve) if n_cve else "")

    posture = ""
    if p:
        def ok(v):
            return '<span class="ok">✓ van</span>' if v else '<span class="no">✗ hiányzik</span>'
        posture = ("<h2>DNS / e-mail posture</h2><ul class=\"post\">"
                   f"<li>SPF: {ok(p.get('spf_ok'))}</li><li>DMARC: {ok(p.get('dmarc_ok'))}</li>"
                   f"<li>DKIM: {E(', '.join(p.get('dkim') or [])) or '✗ nincs'}</li>"
                   f"<li>CAA: {E(', '.join(p.get('caa') or [])) or '✗ nincs'}</li>"
                   f"<li>MX: {E(', '.join(p.get('mx') or [])) or '✗ nincs'}</li>"
                   f"<li>NS: {E(', '.join(p.get('ns') or [])) or '—'}</li></ul>")
        if wc:
            posture += f'<p class="warn">Wildcard DNS: {E(", ".join(wc))}</p>'

    order = {"takeover": 0, "exposed": 1, "live": 2, "dead": 3}
    attn = sorted([h for h in hosts if _status(h) in ("takeover", "exposed")], key=lambda h: (order[_status(h)], h["host"]))
    live = sorted([h for h in hosts if _status(h) == "live"], key=lambda h: h["host"])
    dead = sorted([h for h in hosts if _status(h) == "dead"], key=lambda h: h["host"])
    body = ""
    if attn:
        body += "<h2>⚠ Figyelmet igénylő hostok</h2>" + "".join(_host_html(h) for h in attn)
    if live:
        body += "<h2>Élő hostok</h2>" + "".join(_host_html(h) for h in live)
    if dead:
        body += "<h2>Nem élő / fel nem oldott</h2><p class=\"meta\">" + E(", ".join(h["host"] for h in dead)) + "</p>"

    return f"""<!doctype html><html lang="hu"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><meta name="robots" content="noindex,nofollow">
<title>domain-recon jelentés — {E(domain)}</title>
<link rel="icon" href="/static/favicon.svg" type="image/svg+xml">
<link rel="stylesheet" href="/static/report.css">
<script src="/static/report-print.js" defer></script></head>
<body><div class="page">
<div class="bar no-print"><a href="/#{E(scan.get('id',''))}">← Vissza</a>
<button id="print" type="button">Nyomtatás / PDF mentés</button>
<a href="/api/export/{E(scan.get('id',''))}.md">Markdown</a></div>
<header><div class="brand">SADROBOT // RECON</div>
<h1>Feltérképezési jelentés — {E(domain)}</h1>
<p class="sub">{_fmt_dt(scan.get('created'))} · {ver} · passzív: {E(str(scan.get('status','?')))} · aktív: {E(str(scan.get('active_status','none')))} · CVE: {E(str(scan.get('cve_status','none')))}</p></header>
<div class="stats">{stats}</div>
{posture}
{body}
<footer>sadrobot · domain-recon · csak engedélyezett célpontra</footer>
</div></body></html>"""
