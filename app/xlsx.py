"""Formázott Excel (.xlsx) riport generálása egy domain-recon scanből.

FilterMax márka színek:
  NAVY  (cím, fejléc háttér, tab, host-elválasztó border, alcím) -> mélyzöld 176838
  LIME  (vékony accent sáv a 3. sorban)                          -> lime     8AB046
  PALE  (hostonkénti sávozás)                                    -> halványzöld EAF3DE
A HTTP-status feltételes formázás színei szemantikusak, nem márka -> maradnak.
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.formatting.rule import Rule, FormulaRule, CellIsRule
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

# --- FilterMax márka színek ---
NAVY = "176838"   # mélyzöld (volt navy)
LIME = "8AB046"   # lime (volt lime)
PALE = "EAF3DE"   # halványzöld sávozás

# Oszlop fejléc címkék és szélességek (karakter)
HEADERS = [
    "Host", "Alive", "IP Address(es)", "CNAME", "Port", "Service",
    "HTTP Status", "Server", "Page Title", "TLS CN", "Banner",
]
WIDTHS = [39, 8, 49, 39, 8, 14, 13, 34, 34, 32, 37]  # A..K
# Középre igazított oszlopok: B (Alive), E (Port), G (HTTP Status)
CENTER_COLS = {"B", "E", "G"}


def _build_rows(scan: dict, hosts: list[dict]) -> list[list]:
    """Lapos tábla: hostonként × nyitott portonként egy sor.
    A Host oszlop MINDEN soron ki van töltve (ettől függ a divider + banding).
    Port nélküli host -> egyetlen sor üres port mezőkkel.
    """
    rows = []
    # Hosztokat névre rendezzük, hogy egy host sorai összefüggőek legyenek.
    for h in sorted(hosts, key=lambda x: (x.get("host") or "")):
        host = h.get("host") or ""
        alive = "yes" if h.get("alive") else "no"
        ips = " ".join(h.get("ips") or [])
        cname = h.get("cname") or ""
        ports = h.get("ports") or []
        if not ports:
            rows.append([host, alive, ips, cname, None, "", None, "", "", "", ""])
            continue
        for p in ports:
            http = p.get("http") or {}
            tls = p.get("tls") or {}
            banner = p.get("banner") or ""
            if banner and len(banner) > 120:
                banner = banner[:120]
            port_val = p.get("port")
            status = http.get("status")
            rows.append([
                host,
                alive,
                ips,
                cname,
                int(port_val) if isinstance(port_val, (int, float)) else (port_val if port_val is not None else None),
                p.get("name") or "",
                int(status) if isinstance(status, (int, float)) else (status if status is not None else None),
                http.get("server") or "",
                http.get("title") or "",
                tls.get("cn") or "",
                banner,
            ])
    return rows


def build_bytes(scan: dict, hosts: list[dict], classification: str = "Confidential") -> bytes:
    domain = scan.get("domain") or ""
    scan_id = scan.get("id") or ""
    created = int(scan.get("created") or 0)

    rows = _build_rows(scan, hosts)
    last = 4 + len(rows)  # adat sorok: 5..last

    # Excel dátum serial (Unix epoch -> 1900 rendszer)
    serial = created / 86400 + 25569

    wb = Workbook()
    ws = wb.active
    ws.title = "Domain Recon"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = NAVY

    # --- Row 1: cím + klasszifikáció ---
    c = ws["A1"]
    c.value = f"Domain Recon — {domain}"
    c.font = Font(name="Calibri", size=18, bold=True, color=NAVY)
    k1 = ws["K1"]
    k1.value = classification
    k1.font = Font(name="Calibri", size=10, color="000000")
    k1.alignment = Alignment(horizontal="right")

    # --- Row 2: alcím formula + comment ---
    a2 = ws["A2"]
    serial_str = repr(round(serial, 6))
    a2.value = (
        f'="External attack-surface scan   ·   "&COUNTA($A$5:$A${last})'
        f'&" records   ·   Scan "&TEXT({serial_str},"yyyy-mm-dd hh:mm")'
        f'&"   ·   ID {scan_id}"'
    )
    a2.font = Font(name="Calibri", size=10, color=NAVY)
    a2.comment = Comment(
        f"Source: domain-recon scan\nDomain: {domain}\nScan ID: {scan_id}\n"
        f"Created (unix): {created}\nExcel serial: {serial}",
        "domain-recon",
    )

    # --- Row 3: accent sáv (lime) ---
    accent = PatternFill("solid", fgColor=LIME)
    for col in range(1, 12):  # A..K
        ws.cell(row=3, column=col).fill = accent
    ws.row_dimensions[3].height = 6

    # --- Row 4: fejléc ---
    header_fill = PatternFill("solid", fgColor=NAVY)
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    for i, label in enumerate(HEADERS):
        col = i + 1
        cell = ws.cell(row=4, column=col, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[4].height = 22

    # --- Oszlopszélességek ---
    for i, w in enumerate(WIDTHS):
        ws.column_dimensions[get_column_letter(i + 1)].width = w

    # --- Adat sorok ---
    data_font = Font(name="Calibri", size=11)
    for r_idx, row in enumerate(rows, start=5):
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = data_font
            letter = get_column_letter(c_idx)
            if letter in CENTER_COLS:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(vertical="center")

    # --- Freeze + AutoFilter ---
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:K{last}"

    # --- Feltételes formázás (prioritás: kisebb szám = magasabb prioritás) ---
    # Csak akkor, ha van adat sor (last>4), különben a tartományok érvénytelenek.
    if rows:
        data_range = f"A5:K{last}"
        g_range = f"G5:G{last}"
        prio = 1

        # 1-4. HTTP status színek a G oszlopban
        status_specs = [
            ("200", "299", "D6EEDD", "1F8A4C"),
            ("300", "399", "DCEBF3", "1F5E80"),
            ("400", "499", "FBE6CF", "9A6A12"),
            ("500", "599", "F6D9D5", "B02A1E"),
        ]
        for lo, hi, fill_hex, font_hex in status_specs:
            rule = CellIsRule(
                operator="between",
                formula=[lo, hi],
                fill=PatternFill("solid", fgColor=fill_hex),
                font=Font(color=font_hex),
            )
            rule.priority = prio
            prio += 1
            ws.conditional_formatting.add(g_range, rule)

        # 5. host-elválasztó: felső vékony mélyzöld border, ha $A5<>$A4
        divider_dxf = DifferentialStyle(
            border=Border(top=Side(style="thin", color=NAVY))
        )
        divider_rule = Rule(type="expression", dxf=divider_dxf, stopIfTrue=False)
        divider_rule.formula = ["$A5<>$A4"]
        divider_rule.priority = prio
        prio += 1
        ws.conditional_formatting.add(data_range, divider_rule)

        # 6. hostonkénti sávozás: ISODD(SUMPRODUCT(...)) -> halványzöld
        banding_rule = FormulaRule(
            formula=["ISODD(SUMPRODUCT(--($A$5:$A5<>$A$4:$A4)))"],
            fill=PatternFill("solid", fgColor=PALE),
        )
        banding_rule.priority = prio
        prio += 1
        ws.conditional_formatting.add(data_range, banding_rule)

    # --- Nyomtatás ---
    ws.oddHeader.right.text = classification
    ws.oddFooter.right.text = "&P"

    # --- Kiírás memóriába ---
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
