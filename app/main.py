"""domain-recon – FastAPI. Írj be egy domaint → aldomének, hostok, élő/nem élő, portok, szolgáltatások.

Hitelesítést a fordított proxy (oauth2-proxy/passkey) ad; ez az app egy megbízható felhasználót szolgál ki.
Védelem: TrustedHost, szigorú CSP, kéréstörzs-limit, egyidejű vizsgálatok korlátja, domain-validáció.
"""
from __future__ import annotations

import asyncio
import csv
import hmac
import io
import ipaddress
import json
import os
import re
import secrets
import time
from itertools import islice

from fastapi import Body, FastAPI, HTTPException, Request
from fastapi.responses import FileResponse, JSONResponse, PlainTextResponse, Response
from fastapi.staticfiles import StaticFiles

from . import __version__, engine, report
from .db import DB
from .ports import resolve_profile

HERE = os.path.dirname(__file__)
STATIC = os.path.join(HERE, "static")


def _build_stamp() -> str:
    """Build-bélyeg a telepített verzió követéséhez: RECON_BUILD env vagy app/BUILD fájl."""
    b = os.getenv("RECON_BUILD", "").strip()
    if b:
        return b
    try:
        with open(os.path.join(HERE, "BUILD"), encoding="utf-8") as fh:
            return fh.read().strip()
    except OSError:
        return ""


BUILD = _build_stamp()
DATA = os.getenv("RECON_DATA", os.path.join(os.path.dirname(HERE), "data"))
ALLOWED = [h.strip() for h in os.getenv("RECON_ALLOWED_HOSTS", "").split(",") if h.strip()]
MAX_SCANS = int(os.getenv("RECON_MAX_CONCURRENCY", "2"))
KEEP = int(os.getenv("RECON_KEEP", "60"))
MAX_BODY = 4096
IMPORT_MAX = int(os.getenv("RECON_IMPORT_MAX", str(8 * 1024 * 1024)))   # XLS/CSV-import feltöltési korlát
PROXY_SECRET = os.getenv("RECON_PROXY_SECRET", "")   # ha be van állítva, csak a Caddy X-Recon-Proxy fejlécét fogadjuk el

HOST_RE = re.compile(r"^(?=.{1,253}$)(?!-)[A-Za-z0-9-]{1,63}(?<!-)(\.(?!-)[A-Za-z0-9-]{1,63}(?<!-))+$")

db = DB(os.path.join(DATA, "recon.db"))
app = FastAPI(title="domain-recon", docs_url=None, redoc_url=None, openapi_url=None)

CSP = ("default-src 'none'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; "
       "font-src 'self'; connect-src 'self'; base-uri 'none'; form-action 'none'; frame-ancestors 'none'")
SEC_HEADERS = {
    "Content-Security-Policy": CSP,
    "X-Content-Type-Options": "nosniff",
    "Referrer-Policy": "no-referrer",
    "X-Frame-Options": "DENY",
    "Permissions-Policy": "geolocation=(), microphone=(), camera=()",
}


@app.middleware("http")
async def guard(request: Request, call_next):
    # a /healthz mindig átmegy (belső konténer-healthcheck, tetszőleges Host, proxy-titok nélkül)
    if request.url.path == "/healthz":
        return await call_next(request)
    # mélységi védelem: csak a fordított proxy (Caddy) közös titkát hordozó kérés megy át
    if PROXY_SECRET:
        if not hmac.compare_digest(request.headers.get("x-recon-proxy", ""), PROXY_SECRET):
            return PlainTextResponse("forbidden", status_code=403)
    # Host-fejléc ellenőrzés (ha konfigurálva)
    if ALLOWED:
        host = (request.headers.get("host") or "").split(":")[0]
        if host and host not in ALLOWED:
            return PlainTextResponse("bad host", status_code=400)
    # kéréstörzs-limit
    cl = request.headers.get("content-length")
    limit = IMPORT_MAX if request.url.path == "/api/import" else MAX_BODY   # az XLS/CSV-import nagyobb törzset kaphat
    if cl and cl.isdigit() and int(cl) > limit:
        return PlainTextResponse("payload too large", status_code=413)
    resp = await call_next(request)
    for k, v in SEC_HEADERS.items():
        resp.headers.setdefault(k, v)
    if request.url.path.startswith("/static/"):   # a statikus JS/CSS mindig revalidáljon (ne ragadjon be régi app.js)
        resp.headers["Cache-Control"] = "no-cache"
    return resp


app.mount("/static", StaticFiles(directory=STATIC), name="static")


_NOCACHE = {"Cache-Control": "no-cache"}


def _page(name: str) -> Response:
    """HTML kiszolgálása a verzió behelyettesítésével (asset-URL cache-busting: ?v=<verzió>)."""
    with open(os.path.join(STATIC, name), encoding="utf-8") as fh:
        html = fh.read().replace("__VER__", __version__)
    return Response(html, media_type="text/html; charset=utf-8", headers=_NOCACHE)


@app.get("/")
async def index():
    return _page("index.html")


@app.get("/graph")
async def graph():
    return _page("graph.html")


@app.get("/report")
async def report_page(request: Request):
    q = request.query_params
    sid = q.get("id", "")
    s = db.get_scan(sid)
    if not s:
        raise HTTPException(404, "Nincs ilyen vizsgálat.")
    s["version"], s["build"] = __version__, BUILD
    lang = q.get("lang", "hu")
    focus = q.get("focus", "full")
    cve = q.get("cve", "")[:40]
    return Response(report.build_html(s, db.hosts(sid), lang=lang, focus=focus, cve=cve),
                    media_type="text/html; charset=utf-8", headers=_NOCACHE)


@app.get("/healthz")
async def healthz():
    return PlainTextResponse("ok")


@app.get("/api/version")
async def api_version():
    return {"version": __version__, "build": BUILD}


@app.get("/robots.txt")
async def robots():
    return PlainTextResponse("User-agent: *\nDisallow: /\n")


def _norm_domain(raw: str) -> str:
    d = (raw or "").strip().lower().rstrip(".")[:253]   # hosszkorlát a regex előtt (ReDoS-védelem)
    d = re.sub(r"^https?://", "", d)
    d = d.split("/")[0].split(":")[0]
    try:
        d = d.encode("idna").decode("ascii")
    except Exception:  # noqa: BLE001
        pass
    if not HOST_RE.match(d):
        raise HTTPException(400, "Érvénytelen domain.")
    return d


_MAX_TARGET_IPS = int(os.getenv("RECON_MAX_HOSTS", "750"))


def _expand_ips(s: str):
    """(ips, total). CIDR / tartomány (a-b vagy a.b.c.d-N) / egyetlen IP → IP-lista (capelve) + teljes darabszám."""
    if "/" in s:                                   # CIDR
        net = ipaddress.ip_network(s, strict=False)
        total = net.num_addresses - (2 if net.version == 4 and net.prefixlen < 31 else 0)
        total = max(total, 1)
        ips = [str(x) for x in islice(net.hosts(), _MAX_TARGET_IPS)] or [str(net.network_address)]
        return ips, total
    if "-" in s:                                   # tartomány
        a, _, b = s.partition("-")
        start = ipaddress.ip_address(a.strip())
        if "." in b or ":" in b:
            end = ipaddress.ip_address(b.strip())
        elif start.version == 4 and b.strip().isdigit():
            end = ipaddress.ip_address((int(start) & ~0xFF) | (int(b) & 0xFF))
        else:
            raise ValueError("bad range")
        lo, hi = sorted((int(start), int(end)))
        total = hi - lo + 1
        ips = [str(ipaddress.ip_address(i)) for i in range(lo, min(hi, lo + _MAX_TARGET_IPS - 1) + 1)]
        return ips, total
    ipaddress.ip_address(s)                        # egyetlen IP (ValueError ha nem az)
    return [s], 1


def _parse_targets(cells, allow_private: bool = False) -> tuple[list[str], int]:
    """Cellaszövegekből IP / CIDR / tartomány tokenek → egyedi, szkennelhető IP-lista (capelve) + nyers darabszám.
    A szűrés az engine SSRF-politikájával egyezik (loopback/link-local/reserved/multicast SOHA; privát csak
    allow_private-tal). Nem-IP cellák (fejléc, domain, üres) csendben kimaradnak."""
    seen: set[str] = set()
    out: list[str] = []
    raw_count = 0
    for cell in cells:
        for tok in re.split(r"[\s,;|]+", str(cell if cell is not None else "").strip()):
            if not tok:
                continue
            try:
                ips, _ = _expand_ips(tok)
            except (ValueError, Exception):   # noqa: BLE001  – nem cél: csendben kihagyjuk
                continue
            for ip in ips:
                raw_count += 1
                if engine.scannable(ip, allow_private) and ip not in seen:
                    seen.add(ip)
                    out.append(ip)
                    if len(out) >= _MAX_TARGET_IPS:
                        return out, raw_count
    return out, raw_count


def _xlsx_cells(raw: bytes):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    if v is not None:
                        yield v
    finally:
        wb.close()


def _csv_cells(raw: bytes):
    text = raw.decode("utf-8", "replace")
    for row in csv.reader(io.StringIO(text)):
        for v in row:
            yield v


def _domain_or_none(raw: str):
    """Egy token normalizált domainje, vagy None (nem dob – a _norm_domain hibatűrő párja)."""
    d = (raw or "").strip().lower().rstrip(".")[:253]
    d = re.sub(r"^https?://", "", d).split("/")[0].split(":")[0]
    try:
        d = d.encode("idna").decode("ascii")
    except Exception:  # noqa: BLE001
        pass
    return d if HOST_RE.match(d) else None


def _parse_domains(cells) -> tuple[list[str], int]:
    """Cellaszövegekből domainek (ami NEM IP/CIDR/tartomány) → egyedi, érvényes domain-lista (capelve)."""
    seen: set[str] = set()
    out: list[str] = []
    raw_count = 0
    for cell in cells:
        for tok in re.split(r"[\s,;|]+", str(cell if cell is not None else "").strip()):
            if not tok:
                continue
            try:
                _expand_ips(tok)        # IP/CIDR/tartomány → nem domain (azt a _parse_targets kezeli)
                continue
            except (ValueError, Exception):  # noqa: BLE001
                pass
            d = _domain_or_none(tok)
            if d:
                raw_count += 1
                if d not in seen:
                    seen.add(d)
                    out.append(d)
                    if len(out) >= _MAX_TARGET_IPS:
                        return out, raw_count
    return out, raw_count


@app.post("/api/import")
async def api_import(request: Request):
    """XLS(X)/CSV feltöltés nyers törzsként (?name=fájlnév) → a benne talált, szkennelhető IP-k ÉS domainek listája.
    Nem indít vizsgálatot; a kliens megerősítés után a /api/scan-t hívja (ips és/vagy hosts)."""
    name = (request.query_params.get("name") or "").lower()
    raw = await request.body()
    if len(raw) > IMPORT_MAX:
        raise HTTPException(413, "Túl nagy fájl.")
    if not raw:
        raise HTTPException(400, "Üres fájl.")
    try:
        if name.endswith((".xlsx", ".xlsm")):
            cells = list(_xlsx_cells(raw))
        elif name.endswith((".csv", ".txt", ".tsv")) or raw[:4] != b"PK\x03\x04":
            cells = list(_csv_cells(raw))
        else:                                   # ismeretlen kiterjesztés, de ZIP-fejléc → xlsx
            cells = list(_xlsx_cells(raw))
    except HTTPException:
        raise
    except Exception:  # noqa: BLE001
        raise HTTPException(400, "Nem olvasható fájl (XLSX vagy CSV kell).")
    allow_private = request.query_params.get("private") in ("1", "true", "yes")
    ips, raw_count = _parse_targets(cells, allow_private)
    domains, dom_raw = _parse_domains(cells)
    return {"ips": ips, "domains": domains,
            "count": len(ips), "domain_count": len(domains), "raw": raw_count,
            "truncated": len(ips) >= _MAX_TARGET_IPS or len(domains) >= _MAX_TARGET_IPS,
            "source": request.query_params.get("name") or ""}


def _norm_target(raw: str):
    """(kind, label, ips, total). kind: 'domain' → ips None; 'ip' → cím/CIDR/tartomány."""
    s = (raw or "").strip()
    s = re.sub(r"^https?://", "", s).strip().rstrip(".")
    looks_ip = bool(re.fullmatch(r"[0-9a-fA-F:.\-/]+", s)) and (":" in s or "/" in s or bool(re.search(r"\d+\.\d+", s)))
    if looks_ip:
        try:
            ips, total = _expand_ips(s)
            if ips:
                return "ip", s, ips[:_MAX_TARGET_IPS], total
        except ValueError:
            pass
    return "domain", _norm_domain(raw), None, 0


def _active_opts(payload: dict) -> dict:
    ports = str(payload.get("ports", "service"))
    custom = str(payload.get("custom_ports", ""))[:2000]
    _, pname = resolve_profile(ports, custom)
    return {
        "ports": ports, "custom_ports": custom, "profile": pname,
        "intensity": payload.get("intensity", "polite") if payload.get("intensity") in engine.INTENSITY else "polite",
        "scan_ports": bool(payload.get("scan_ports", True)),
        "brute": bool(payload.get("brute", False)),
        "allow_private": bool(payload.get("allow_private", False)),
    }


@app.post("/api/scan")
async def api_scan(payload: dict = Body(...)):
    running = sum(1 for s in db.recent_scans(MAX_SCANS * 3)
                  if s["status"] in ("queued", "running") or s.get("active_status") in ("queued", "running"))
    if running >= MAX_SCANS:
        raise HTTPException(429, f"Már {running} vizsgálat fut. Várd meg, amíg befejeződik.")
    hosts_in = payload.get("hosts")
    ips_in = payload.get("ips")
    opts = _active_opts(payload)
    src = str(payload.get("source") or "").strip()[:60]
    if isinstance(hosts_in, list) and hosts_in:          # XLS/CSV-import: domain-lista (leltár; IP-k is mehetnek)
        doms, _ = _parse_domains(hosts_in)
        ips2, _ = _parse_targets(ips_in or [], bool(payload.get("allow_private")))
        if not doms and not ips2:
            raise HTTPException(400, "Nincs érvényes domain/IP a listában.")
        total = len(doms) + len(ips2)
        label = (f"{total} célpont" + (f" · {src}" if src else ""))[:120]
        opts["target_kind"] = "hostlist"
        opts["target_hosts"] = doms
        opts["target_ips"] = ips2
        opts["target_total"] = total
    elif isinstance(ips_in, list) and ips_in:            # XLS/CSV-import: kész IP-lista
        ips, _raw = _parse_targets(ips_in, bool(payload.get("allow_private")))
        if not ips:
            raise HTTPException(400, "Nincs érvényes (publikus) IP a listában.")
        total = len(ips)
        label = (f"{total} IP" + (f" · {src}" if src else ""))[:120]
        opts["target_kind"] = "ip"
        opts["target_ips"] = ips
        opts["target_total"] = total
    else:
        kind, label, ips, total = _norm_target(payload.get("domain", ""))
        opts["target_kind"] = kind
        if kind == "ip":
            opts["target_ips"] = ips
            opts["target_total"] = total
    scan_id = secrets.token_hex(8)
    db.create_scan(scan_id, label, opts)
    db.prune(KEEP)
    engine.start_passive(db, scan_id, label, opts)
    return {"id": scan_id, "domain": label, "opts": opts}


@app.post("/api/scan/{scan_id}/active")
async def api_scan_active(scan_id: str, payload: dict = Body(default={})):
    s = db.get_scan(scan_id)
    if not s:
        raise HTTPException(404, "Nincs ilyen vizsgálat.")
    if s["status"] in ("queued", "running"):
        raise HTTPException(409, "A passzív fázis még fut.")
    if s.get("active_status") in ("queued", "running"):
        raise HTTPException(409, "Az aktív fázis már fut.")
    opts = _active_opts({**(s.get("opts") or {}), **(payload or {})})
    engine.start_active(db, scan_id, opts)
    return {"id": scan_id, "active": "running", "opts": opts}


@app.post("/api/scan/{scan_id}/cve")
async def api_scan_cve(scan_id: str):
    s = db.get_scan(scan_id)
    if not s:
        raise HTTPException(404, "Nincs ilyen vizsgálat.")
    if s["status"] in ("queued", "running"):
        raise HTTPException(409, "A passzív fázis még fut.")
    if s.get("cve_status") in ("queued", "running"):
        raise HTTPException(409, "A CVE-vizsgálat már fut.")
    engine.start_cve(db, scan_id)
    return {"id": scan_id, "cve": "running"}


@app.post("/api/scan/{scan_id}/subdomains")
async def api_scan_subdomains(scan_id: str):
    """Domain-leltár külön fázisa: a betöltött domainekre CT-alapú aldomén-felderítés."""
    s = db.get_scan(scan_id)
    if not s:
        raise HTTPException(404, "Nincs ilyen vizsgálat.")
    if (s.get("opts") or {}).get("target_kind") != "hostlist":
        raise HTTPException(400, "Aldomén-felderítés csak domain-leltárra indítható.")
    if s["status"] in ("queued", "running"):
        raise HTTPException(409, "A leltár még fut.")
    if s.get("sub_status") in ("queued", "running"):
        raise HTTPException(409, "Az aldomén-felderítés már fut.")
    engine.start_subdomains(db, scan_id)
    return {"id": scan_id, "subdomains": "running"}


@app.get("/api/scans")
async def api_scans():
    return {"scans": db.recent_scans(40)}


@app.get("/api/scan/{scan_id}")
async def api_scan_get(scan_id: str):
    s = db.get_scan(scan_id)
    if not s:
        raise HTTPException(404, "Nincs ilyen vizsgálat.")
    s["hosts"] = db.hosts(scan_id)
    s["now"] = int(time.time())
    return JSONResponse(s)


@app.post("/api/scan/{scan_id}/cancel")
async def api_scan_cancel(scan_id: str):
    if not db.get_scan(scan_id):
        raise HTTPException(404, "Nincs ilyen vizsgálat.")
    ok = engine.cancel_scan(scan_id)
    return {"canceled": ok}


@app.delete("/api/scan/{scan_id}")
async def api_scan_delete(scan_id: str):
    if not db.get_scan(scan_id):
        raise HTTPException(404, "Nincs ilyen vizsgálat.")
    engine.cancel_scan(scan_id)
    db.delete_scan(scan_id)
    return {"deleted": True}


def _scan_for_export(scan_id: str):
    s = db.get_scan(scan_id)
    if not s:
        raise HTTPException(404, "Nincs ilyen vizsgálat.")
    s["version"], s["build"] = __version__, BUILD
    return s, db.hosts(scan_id)


@app.get("/api/export/{scan_id}.json")
async def api_export_json(scan_id: str):
    s, hosts = _scan_for_export(scan_id)
    data = json.dumps(report.build_json(s, hosts), ensure_ascii=False, indent=2)
    return Response(data, media_type="application/json",
                    headers={"Content-Disposition": f'attachment; filename="recon_{s["domain"]}_{scan_id}.json"'})


@app.get("/api/export/{scan_id}.md")
async def api_export_md(scan_id: str):
    s, hosts = _scan_for_export(scan_id)
    data = report.build_markdown(s, hosts)
    return Response(data, media_type="text/markdown; charset=utf-8",
                    headers={"Content-Disposition": f'attachment; filename="recon_{s["domain"]}_{scan_id}.md"'})


@app.get("/api/export/{scan_id}.pdf")
async def api_export_pdf(scan_id: str, request: Request):
    s, hosts = _scan_for_export(scan_id)
    q = request.query_params
    lang = q.get("lang", "hu")
    focus = q.get("focus", "full")
    cve = q.get("cve", "")[:40]
    data = report.build_pdf(s, hosts, lang=lang, focus=focus, cve=cve)
    suffix = (f"_{focus}" if focus != "full" else "") + (f"_{cve}" if cve else "")
    return Response(data, media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="recon_{s["domain"]}_{scan_id}{suffix}.pdf"'})


@app.get("/api/export/{scan_id}.xlsx")
async def api_export_xlsx(scan_id: str):
    s, hosts = _scan_for_export(scan_id)
    data = report.build_xlsx(s, hosts)
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="recon_{s["domain"]}_{scan_id}.xlsx"'})


def _csv_safe(v) -> str:
    """CSV-képletinjekció ellen: a =,+,-,@ (és vezető tab/CR) kezdetű cellát aposztróffal semlegesítjük."""
    s = "" if v is None else str(v)
    if s and s[0] in "=+-@\t\r":
        return "'" + s
    return s


@app.get("/api/export/{scan_id}.csv")
async def api_export_csv(scan_id: str):
    s, hosts = _scan_for_export(scan_id)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["# domain-recon", s["domain"], scan_id, time.strftime("%Y-%m-%d %H:%M", time.localtime(s["created"]))])
    cols = ["host", "status", "ips", "ptr", "asn", "org", "cc", "cname",
            "port", "service", "http_status", "server", "title", "tech", "tls_cn", "banner", "findings"]
    w.writerow(cols)
    first = True
    for h in sorted(hosts, key=lambda x: (report._status(x) != "takeover", report._status(x) != "exposed", not x.get("alive"), x["host"])):
        if not first:
            w.writerow([])                      # üres sor választja el a hostokat
        first = False
        info0 = (h.get("ips_info") or [{}])[0]
        findings = "; ".join([f"{v.get('severity')}:{v.get('title')}" for v in (h.get("vuln") or []) + (h.get("cve") or [])]
                             + (h.get("exposed_flags") or [])
                             + ([f"takeover:{h['takeover']['service']}"] if h.get("takeover", {}).get("service") else []))
        hostcells = [h["host"], report._status(h), " ".join(h.get("ips", [])), info0.get("ptr") or "",
                     (f"AS{info0['asn']}" if info0.get("asn") else ""), info0.get("org") or "", info0.get("cc") or "",
                     h.get("cname") or ""]
        ports = h.get("ports") or []
        if not ports:
            w.writerow([_csv_safe(x) for x in hostcells + ["", "", "", "", "", "", "", "", findings]])
            continue
        for idx, p in enumerate(ports):
            http = p.get("http") or {}
            tls = p.get("tls") or {}
            lead = hostcells if idx == 0 else ["", "", "", "", "", "", "", ""]   # host-mezők csak az első sorban
            w.writerow([_csv_safe(x) for x in lead + [
                p["port"], p.get("name", ""), http.get("status", ""), http.get("server", ""),
                http.get("title", ""), ", ".join(p.get("tech", [])), tls.get("cn", ""),
                (p.get("banner") or "")[:120], findings if idx == 0 else ""]])
    return Response(buf.getvalue(), media_type="text/csv",
                    headers={"Content-Disposition": f'attachment; filename="recon_{s["domain"]}_{scan_id}.csv"'})
